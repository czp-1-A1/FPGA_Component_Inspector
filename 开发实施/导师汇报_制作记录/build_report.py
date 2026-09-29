from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from pathlib import Path
D=Document(); sec=D.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.65); sec.bottom_margin=Cm(1.55); sec.left_margin=sec.right_margin=Cm(1.8)
for name in ['Normal','Title','Heading 1','Heading 2']:
 s=D.styles[name]; s.font.name='Microsoft YaHei'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'Microsoft YaHei'); s.font.color.rgb=RGBColor(0,0,0)
s=D.styles['Normal']; s.font.size=Pt(10.5); s.paragraph_format.line_spacing=1.12; s.paragraph_format.space_after=Pt(5)
for name,size in [('Title',19),('Heading 1',14),('Heading 2',11.5)]:
 s=D.styles[name]; s.font.size=Pt(size); s.font.bold=True; s.paragraph_format.space_before=Pt(9); s.paragraph_format.space_after=Pt(6)
D.core_properties.title='FPGA载带视觉检测项目方向与可行性说明'; D.core_properties.subject='导师汇报'; D.core_properties.author=''
def p(t,b=False,size=None):
 x=D.add_paragraph(); r=x.add_run(t); r.bold=b
 if size:r.font.size=Pt(size)
 return x
def h(t):D.add_heading(t,1)
def table(headers,rows,widths):
 t=D.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
 pr=t._tbl.tblPr; borders=OxmlElement('w:tblBorders')
 for e in ['top','left','bottom','right','insideH','insideV']:
  x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'D9D9D9'); borders.append(x)
 pr.append(borders)
 for i,w in enumerate(widths):t.columns[i].width=Cm(w)
 for n,vals in enumerate([headers]+rows):
  cells=t.rows[0].cells if n==0 else t.add_row().cells
  for i,txt in enumerate(vals):
   c=cells[i]; c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   tcpr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
   for side in ['top','bottom','left','right']:
    e=OxmlElement('w:'+side); e.set(qn('w:w'),'65'); e.set(qn('w:type'),'dxa'); mar.append(e)
   tcpr.append(mar)
   if n==0:
    sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'E8EFF5'); tcpr.append(sh)
   para=c.paragraphs[0]; para.paragraph_format.space_after=Pt(0); para.paragraph_format.line_spacing=1.08
   r=para.add_run(txt); r.font.size=Pt(9.5); r.bold=n==0
  trpr=t.rows[n]._tr.get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
  if n==0:trpr.append(OxmlElement('w:tblHeader'))
 D.add_paragraph().paragraph_format.space_after=Pt(0)
 return t
D.add_heading('FPGA载带视觉检测项目\n方向与可行性说明',0)
p('项目名称  基于安路FPGA的电子元件载带视觉检测系统',True)
p('导师交流稿  |  2026年9月17日  |  计划周期约六周',size=9)
p('拟用摄像头观察载带中的元件，由FPGA实时判断缺件、位置或方向异常，在显示器上标注结果并记录异常穴位编号。整条检测结束后，按清单交由人工维修。第一版取消精准停车。',True)
h('一  做什么与怎样工作')
p('载带可理解为装元件的连续“小格子”。本项目检查每个格子中是否有料、是否放正，并尝试依据顶部标记识别装反。首版限定一种元件和一种载带。')
x=p('载带送料 → 摄像头采集 → FPGA局部图像检测\n↓\nHDMI显示结果与缺陷清单 → 整条交接 → 人工定位维修',True); x.alignment=WD_ALIGN_PARAGRAPH.CENTER
h('二  检测功能分层')
table(['功能','大致判断方法','实现条件与安排'],[
['空穴与缺件','穴内亮度、面积或模板差异','首项必做；空穴和有料必须可稳定区分'],
['明显偏位','元件中心与正常位置比较','补偿载带偏移后验证，优先候选'],
['明显旋转与歪斜','长边方向或区域分布','外形具有方向特征，优先候选'],
['极性装反','色带、圆点或缺口所在一侧','顶部标记清晰可见时加入'],
['翻面与侧立','外形、面积及表面差异','与正常外观差异明显时探索'],
['明显错料','尺寸、颜色、外形或标记','仅针对外观可区分的已知样品'],
['明显破损与异物','轮廓或异常区域','缺陷足够大且无遮挡时探索']],[2.6,5.4,9.4])
p('极性检测可以加入，但检测的是外观标记所代表的装载方向，不是电气测量。无外观差异的错料、阻容值、电性能和隐藏缺陷不属于本阶段承诺。',size=10)
p('第一版目标：空穴检测＋实时显示＋安路Logo；成像验证后，再从偏位、旋转、极性中加入一种最稳定的检测。上述候选均不代表已经实现。',True)
D.add_page_break()
h('三  板卡资源是否足够')
p('目标平台为赛题指定的HX1P35A开发板，FPGA为PH1P35MDG324。官方Lab 1已提供MIPI、图像处理前端、DDR帧缓存和HDMI视频底座。[1][2]')
table(['资源','总量','Lab 1占用','算术剩余'],[
['Slice逻辑单元','21,216','12,187（57.44%）','9,029'],['寄存器','42,432','13,318（31.39%）','29,114'],['ERAM存储块','108','31（28.70%）','77'],['DSP计算单元','40','8（20%）','32']],[4.1,3.4,6.2,3.7])
p('判断：有依据继续使用现有板卡做轻量检测，暂不需要换板。剩余数量不等于所有功能都能加入；完整设计须重新综合、布局布线和上板验证。',True)
p('实现策略：沿用Lab 1视频底座，检测只计算局部图像区域的面积、位置及标记差异；共享计算单元，用ERAM存储字模与缺陷记录，制作轻量结果叠加，保留安路Logo。避免直接堆叠全部官方例程。')
p('时序风险：Lab 1、Lab 2边缘检测、Lab 3完整OSD的最差建立余量分别为0.629、0.234、0.121 ns。它们是不同官方工程，不能视为本项目结果；时序覆盖率分别为95.64%、94.31%、96.29%，后续还需核对约束。[2] ',size=10)
h('四  是否符合赛题方向')
table(['赛题要求','本项目拟对应实现'],[
['基础至少完成四项中的三项','按四项安排：MIPI采集与HDMI；FPGA图像处理；稳定视频；图形文字及安路Logo'],
['FPGA处理并输出处理后视频','提供可观察的灰度或二值化结果画面，避免仅内部计算而无处理结果展示'],
['扩展至少完成一项','目标检测与识别，并在画面中叠加结果；缺陷清单体现实际用途'],
['最终设计部署要求','所有展示功能集成在同一设计中，固化后上电运行，无需PC再次下载']],[5.2,12.2])
p('结论：方案与赛题方向匹配，取消精准停车不影响按目标检测扩展方向开展；当前尚不能写成已满足最终验收要求。[1]',True)
h('五  当前进度与完成把握')
p('截至本次沟通，尚未跑通官方视频，实物检测特征待验证，准备时间约一个半月。固定样品空穴检测有较好的实现基础；偏位和极性有条件可行；连续编号可靠性是主要集成难点。没有足够证据给出可信的成功率百分比，也不承诺通用识别或工业级准确率。',size=10)
D.add_page_break()
h('六  六周实施路线')
table(['时间','主要工作','阶段检查点'],[
['第1周','视频与成像验证','官方视频、固化冷启动；毫米尺视场和正常／异常样品清晰度'],
['第2周','静态空穴检测','先用图片验证特征，再移植FPGA；显示检测区域、结果与Logo'],
['第3周','增加一种检测','选偏位、旋转或极性中最稳定者；保存误检漏检及资源时序结果'],
['第4周','连续检测与编号','明确起点与方向；避免重复计数；编号失可信时标记无效并重检'],
['第5周','缺陷清单与维修交接','穴位号、类型、有效状态可分页查看；满容量提示且不静默覆盖'],
['第6周','联合验证与演示','整条重复测试、重新检测、冷启动；记录错误率与实际处理速度']],[1.6,4.4,11.4])
p('收缩规则：若前两周视频或成像仍不稳定，先完成空穴检测与可靠显示，暂停新增缺陷种类。',True,size=10)
h('七  最大问题与验证办法')
p('首要风险是看不清。用正常件、空穴、偏位及反向样品重复拍摄，先确定工作距离、视场、光照与标记清晰度；不可见的极性特征无法靠算法补救。')
p('第二风险是编号对不上实物。人工标记载带起点和方向，以此开始计数；漏计、滑移或视频中断导致失锁时，标明失效区段并重新检测。维修交接保留起点、方向及清单，编号不作为整卷永久绝对地址。')
p('验证证据：制作有真值编号的样品，调参和验证样品分开；分别统计正常件误报、各类缺陷漏检和编号错误。检查电机启停干扰、清单容量及冷启动。每加一模块保存资源和最终时序报告。机械模型、打印小样和官方例程均不能替代整机实测。',size=10)
h('八  导师意见')
p('建议优先验证的缺陷类型：____________________________________',size=10)
p('方向与范围调整建议：________________________________________',size=10)
p('__________________________________________________________',size=10)
p('资料依据',True,9)
p('[1] 项目目录“赛题要求与分析”内《FPGA赛题解析》资料包：第5页基础要求，第6页扩展要求，第8页平台，第26页固化及集成要求（对应content.md的Slide编号）。',size=8)
p('[2] 项目Example目录下三个官方压缩包：01_MIPI采集与HDMI显示/lab_hd_1_mipi_hdmi.zip；02_边缘检测/lab_hd_2_edge.zip.zip；03_OSD叠加/lab_hd_3_osd.zip。各包内工程根目录/td_project/camera_to_dsi_display_Runs/phy_1/下的camera_to_dsi_display_phy.area与final_timing.rpt。表中“剩余”为总量减官方Lab 1占用。',size=8)
# Page number
f=sec.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=f.add_run(); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); r._r.addnext(fld)
out=Path('E:/2026FPGA/FPGA_Component_Inspector/开发实施/FPGA载带视觉检测项目方向与可行性说明_导师版_20260917.docx')
for el in list(D.styles.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
for el in list(D.element.iter(qn('w:pBdr'))): el.getparent().remove(el)
D.save(out); print(out)

